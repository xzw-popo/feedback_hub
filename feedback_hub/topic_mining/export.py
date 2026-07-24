"""Deterministic and formula-safe exports for already verified topic runs."""
from __future__ import annotations

import json
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Mapping, Sequence

from .run_store import TopicRunStore
from .protocol import classification_protocol
from .service import RunVerificationError, _effective_data_cutoff, _load_manifest, _publish_terminal_mutation, _recall_from_dict, _require_run, _result_scope, _validate_final_rows, _validate_run_identity, _verify_manifest, read_verified_artifact_bytes
from .contracts import load_persisted_topic_spec
from feedback_hub.jsonl_io import load_jsonl_objects


HEADERS = [
    "反馈时间", "反馈原文", "对应链接", "平台", "版本", "设备",
    "Feedback ID", "判定理由", "证据",
]


def export_topic_run(run_id: str, export_format: str, *, store: TopicRunStore | None = None) -> Path:
    if export_format not in {"xlsx", "jsonl"}:
        raise ValueError("unsupported export format")
    if store is None:
        from .service import default_store
        store = default_store()
    run = _require_run(run_id, store)
    if run["status"] != "verified":
        raise RunVerificationError("run_not_verified")
    artifact_dir = Path(run["artifact_dir"])
    spec = load_persisted_topic_spec(json.loads(run["spec_json"]))
    final_path = artifact_dir / "final_reviewed.jsonl"
    manifest = _load_manifest(run, artifact_dir)
    _verify_manifest(manifest, artifact_dir)
    data_cutoff_ms = _effective_data_cutoff(
        run, manifest, require_snapshot=True,
    )
    _validate_run_identity(run, spec, data_cutoff_ms)
    try:
        rows = load_jsonl_objects(
            read_verified_artifact_bytes(manifest, final_path),
        )
    except (UnicodeDecodeError, json.JSONDecodeError, TypeError, ValueError) as exc:
        raise RunVerificationError("invalid_artifact") from exc
    _validate_final_rows(
        rows, spec, expected_run_id=run_id,
        expected_data_cutoff_ms=data_cutoff_ms,
    )
    rows = sorted(rows, key=lambda row: (str(row["label"]), -int(row["source_item"].get("ts_ms", 0)), str(row["item_id"])))
    recall_by_id = _verified_recall_by_id(manifest, artifact_dir)
    if recall_by_id and not {str(row["item_id"]) for row in rows} <= set(recall_by_id):
        raise RunVerificationError("final_result_not_in_recall_pool")
    scope = _result_scope(spec, rows, recall_by_id, manifest)
    selected_rows = list(rows)
    if len(selected_rows) != scope["returned_feedback"]:
        raise RunVerificationError("result_scope_mismatch")
    export_rows = [_export_row(row) for row in selected_rows]
    required_fields = tuple(spec.output["required_fields"])
    if any(
        not set(required_fields).issubset(export_row)
        for export_row in export_rows
    ):
        raise RunVerificationError("required_output_fields_missing")
    jsonl_path = artifact_dir / "final_results.jsonl"
    jsonl_bytes = "".join(
        json.dumps(dict(row), ensure_ascii=False, sort_keys=True) + "\n"
        for row in export_rows
    ).encode("utf-8")
    report = {
        "run_id": run_id, "status": "verified", "matched_count": len(rows),
        "artifact": jsonl_path.name, "data_cutoff_ms": data_cutoff_ms,
        "data_cutoff_time": _format_time(data_cutoff_ms),
        "required_fields": list(required_fields),
        **scope,
    }
    protocol_version, protocol_owner = classification_protocol(run)
    if protocol_version in {2, 3} and protocol_owner == "caller_ai":
        report.update({
            "classification_protocol_version": protocol_version,
            "classification_owner": protocol_owner,
            "classification_revision_count": manifest.get(
                "classification_revision_count", 0,
            ),
            "pending_decision_count": manifest.get(
                "pending_decision_count", 0,
            ),
        })
    report_bytes = (
        json.dumps(report, ensure_ascii=False, sort_keys=True, indent=2) + "\n"
    ).encode("utf-8")
    files = {
        jsonl_path.name: jsonl_bytes,
        "quality_report.json": report_bytes,
    }
    if export_format == "jsonl":
        manifest.update(scope)
        _publish_terminal_mutation(
            run, store, manifest, stage="verified", status="verified",
            files=files,
        )
        return jsonl_path
    workbook_path = artifact_dir / "feedback_list.xlsx"
    files[workbook_path.name] = _xlsx_bytes(selected_rows, scope)
    manifest.update(scope)
    _publish_terminal_mutation(
        run, store, manifest, stage="verified", status="verified",
        files=files,
    )
    return workbook_path


def _verified_recall_by_id(
    manifest: Mapping[str, Any], artifact_dir: Path,
) -> dict[str, Any]:
    recall_path = artifact_dir / "recall_candidates.jsonl"
    artifacts = manifest.get("artifacts")
    if not recall_path.is_file() and (
        not isinstance(artifacts, Mapping) or recall_path.name not in artifacts
    ):
        return {}
    try:
        recalls = [
            _recall_from_dict(row)
            for row in load_jsonl_objects(
                read_verified_artifact_bytes(manifest, recall_path),
            )
        ]
    except (UnicodeDecodeError, json.JSONDecodeError, TypeError, ValueError) as exc:
        raise RunVerificationError("invalid_artifact") from exc
    recall_by_id = {recall.item_id: recall for recall in recalls}
    if len(recall_by_id) != len(recalls):
        raise RunVerificationError("duplicate_item_id")
    return recall_by_id


def _write_xlsx(
    path: Path, rows: Sequence[Mapping[str, Any]], scope: Mapping[str, Any],
) -> None:
    path.write_bytes(_xlsx_bytes(rows, scope))


def _export_row(row: Mapping[str, Any]) -> dict[str, Any]:
    """Add the stable public field names promised by the topic spec."""
    item = row["source_item"]
    return {
        **dict(row),
        "feedback_text": item["text"],
        "feedback_time": _format_time(item["ts_ms"]),
        "source_url": item["source_url"],
    }


def _xlsx_bytes(
    rows: Sequence[Mapping[str, Any]], scope: Mapping[str, Any],
) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "反馈清单"
    sheet.append(HEADERS)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        item = row["source_item"]
        url = str(item["source_url"])
        values = [
            _format_time(item.get("ts_ms")), item["text"], _safe_cell(url),
            item.get("platform", ""), item.get("appversion", ""),
            item.get("device_name", ""), item.get("feedback_id", ""),
            row["reason"], "\n".join(row["evidence"]),
        ]
        sheet.append([_safe_cell(value) for value in values])
        link = sheet.cell(sheet.max_row, 3)
        formula_url = url.replace('"', '""')
        link.value = f'=HYPERLINK("{formula_url}","打开反馈")'
        link.hyperlink = url
        link.style = "Hyperlink"
    sheet.freeze_panes = "A2"
    last_column = openpyxl.utils.get_column_letter(len(HEADERS))
    sheet.auto_filter.ref = f"A1:{last_column}{max(sheet.max_row, 1)}"
    widths = [20, 48, 42, 12, 16, 20, 18, 28, 32]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    metadata = workbook.create_sheet("导出元数据")
    for key in (
        "mode", "result_scope", "matched_total", "returned_feedback",
        "possibly_more_matches",
    ):
        metadata.append([key, scope[key]])
    metadata.column_dimensions["A"].width = 26
    metadata.column_dimensions["B"].width = 24
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _safe_cell(value: Any) -> str:
    text = "" if value is None else str(value)
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def _format_time(value: Any) -> str:
    try:
        return datetime.fromtimestamp(int(value) / 1000, tz=timezone.utc).isoformat()
    except (TypeError, ValueError, OSError):
        return ""
