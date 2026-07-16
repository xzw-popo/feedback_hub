"""Select daily insights and render Markdown and XLSX review artifacts."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from feedback_hub.topic_discovery.global_selection import enrich_insights_with_evidence


_TYPE_ORDER = {
    "rising_or_repeated_bug": 0,
    "new_bug": 1,
    "demand_opportunity": 2,
    "high_value_single": 3,
    "not_reportable": 4,
}

_TYPE_LABELS = {
    "rising_or_repeated_bug": "重复或趋势问题",
    "new_bug": "新问题信号",
    "demand_opportunity": "需求机会",
    "high_value_single": "高价值单条",
    "not_reportable": "不进入报告",
}


def partition_selected_insights(
    insights: list[dict[str, Any]],
    candidates: dict[str, dict[str, Any]],
    global_selection: dict[str, Any],
    *,
    shortlist: list[dict[str, Any]],
) -> dict[str, list[dict[str, Any]]]:
    """Partition local insights using only a validated global ranked selection."""
    normalized = enrich_insights_with_evidence(insights, candidates)
    by_id = {str(row["insight_id"]): row for row in normalized}
    shortlist_by_id = {str(row.get("insight_id") or ""): row for row in shortlist}
    if not all(shortlist_by_id) or len(shortlist_by_id) != len(shortlist):
        raise ValueError("shortlist insight_id must be present and unique")
    if any(insight_id not in by_id for insight_id in shortlist_by_id):
        raise ValueError("shortlist references an unknown insight")
    for insight_id, shortlist_row in shortlist_by_id.items():
        by_id[insight_id].update({
            "shortlist_lanes": list(shortlist_row.get("shortlist_lanes") or []),
            "shortlist_lane_rank": shortlist_row.get("shortlist_lane_rank"),
            "shortlist_rank_fields": dict(shortlist_row.get("shortlist_rank_fields") or {}),
        })

    selections = list(global_selection.get("selected_insights") or [])
    if len(selections) > 5:
        raise ValueError("global selection cannot exceed five insights")
    selected_ids = [str(row.get("insight_id") or "") for row in selections]
    if len(selected_ids) != len(set(selected_ids)):
        raise ValueError("global selection contains duplicate insight IDs")
    if any(insight_id not in by_id for insight_id in selected_ids):
        raise ValueError("global selection references an unknown insight")
    if any(insight_id not in shortlist_by_id for insight_id in selected_ids):
        raise ValueError("global selection references a non-shortlisted insight")
    ordered = sorted(selections, key=lambda row: int(row.get("rank") or 0))
    if [int(row.get("rank") or 0) for row in ordered] != list(range(1, len(ordered) + 1)):
        raise ValueError("global selection ranks must be consecutive")

    main = []
    for selection in ordered:
        insight_id = str(selection["insight_id"])
        source = by_id[insight_id]
        if source.get("needs_human_review"):
            raise ValueError("selected insight requires human review")
        if not source.get("source_links"):
            raise ValueError("selected insight requires a source link")
        row = dict(source)
        row.update({
            "report_decision": "main",
            "global_rank": int(selection["rank"]),
            "global_selection_reason": str(selection.get("selection_reason") or "").strip(),
            "global_report_summary": str(selection.get("report_summary") or "").strip(),
            "editorial_note": str(selection.get("editorial_note") or "").strip(),
        })
        main.append(row)

    selected_id_set = set(selected_ids)
    observe = []
    for source in normalized:
        insight_id = str(source["insight_id"])
        if insight_id in selected_id_set or source.get("report_decision") in {"exclude", "manual_review"}:
            continue
        row = dict(source)
        row["report_decision"] = "observe"
        row["artifact_reason"] = (
            "not_selected_globally" if insight_id in shortlist_by_id else "not_shortlisted"
        )
        observe.append(row)

    return {
        "main": main,
        "observe": sorted(observe, key=_insight_sort_key),
        "manual_review": sorted(
            (row for row in normalized if row.get("report_decision") == "manual_review"),
            key=_insight_sort_key,
        ),
        "exclude": sorted(
            (row for row in normalized if row.get("report_decision") == "exclude"),
            key=_insight_sort_key,
        ),
    }


def render_daily_report(
    report_date: str,
    selected: dict[str, list[dict[str, Any]]],
    media_appendix: list[dict[str, Any]],
    run_summary: dict[str, Any],
) -> str:
    """Render a concise diagnostic daily report draft."""
    baseline_start = str(run_summary.get("baseline_start") or "")
    baseline_end = str(run_summary.get("baseline_end") or "")
    lines = [
        f"# 微信输入法用户反馈日报草稿（{report_date}）",
        "",
        f"数据日：{report_date}；历史基线：{baseline_start} 至 {baseline_end}。频次均按去重会话计算。",
        "",
        "## 今日重点",
        "",
    ]
    main = list(selected.get("main") or [])
    if not main:
        lines.append("本日没有证据充分、值得主动推送的重点洞察。")
    for index, insight in enumerate(main, 1):
        lines.extend(_render_insight(index, insight))

    lines.extend(["", "## 继续观察", ""])
    observe = list(selected.get("observe") or [])[:5]
    if not observe:
        lines.append("暂无。")
    for index, insight in enumerate(observe, 1):
        lines.extend(_render_insight(index, insight, compact=True))

    lines.extend(["", "## 媒体附录", ""])
    if not media_appendix:
        lines.append("暂无仅依赖图片或视频才能判断的反馈。")
    for index, row in enumerate(media_appendix, 1):
        summary = str(row.get("summary") or row.get("title") or "文本信息不足，需查看媒体").strip()
        links = _media_links(row)
        lines.append(f"{index}. {summary}")
        for link in links[:3]:
            lines.append(f"   - [查看原反馈]({link})")

    lines.extend([
        "",
        "## 运行说明",
        "",
        f"模型渠道调用分布：{_json_text(run_summary.get('route_counts') or {})}。",
        "本报告是离线实验草稿；主题分组仅用于报告表达，不修改动态主题库。",
        "",
    ])
    return "\n".join(lines)


def write_review_workbook(
    path: str | Path,
    *,
    selected: dict[str, list[dict[str, Any]]],
    candidates: list[dict[str, Any]],
    relations: list[dict[str, Any]],
    media_appendix: list[dict[str, Any]],
) -> None:
    """Write the four-sheet human review workbook."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    main_sheet = workbook.active
    main_sheet.title = "正文与观察"
    _write_main_sheet(main_sheet, selected)
    _write_candidate_sheet(workbook.create_sheet("全部候选"), candidates)
    _write_relation_sheet(workbook.create_sheet("原子主题关系"), relations)
    _write_media_sheet(workbook.create_sheet("媒体附录"), media_appendix)
    for sheet in workbook.worksheets:
        _format_sheet(sheet)
    workbook.save(path)


def _insight_sort_key(row: dict[str, Any]) -> tuple[int, float, str]:
    return (
        _TYPE_ORDER.get(str(row.get("signal_type") or "not_reportable"), 99),
        -float(row.get("confidence") or 0.0),
        str(row.get("insight_id") or ""),
    )


def _render_insight(index: int, insight: dict[str, Any], *, compact: bool = False) -> list[str]:
    label = _TYPE_LABELS.get(str(insight.get("signal_type") or ""), "观察信号")
    lines = [f"### {index}. {insight.get('headline')}", ""]
    lines.append(f"- 类型：{label}")
    lines.append(f"- 判断：{insight.get('global_report_summary') or insight.get('summary')}")
    lines.append(
        f"- 依据：{insight.get('global_selection_reason') or insight.get('selection_reason')}"
    )
    lines.append(
        f"- 频次：当日 {int(insight.get('today_conversation_count') or 0)} 个去重会话；"
        f"基线逐日 {_json_text(insight.get('baseline_daily_counts') or [])}"
    )
    if not compact:
        lines.append(f"- 趋势口径：{insight.get('trend_claim') or 'none'}")
    for link in (insight.get("source_links") or [])[:3]:
        lines.append(f"- [查看原反馈]({link})")
    lines.append("")
    return lines


def _write_main_sheet(sheet, selected: dict[str, list[dict[str, Any]]]) -> None:
    headers = [
        "report_layer", "insight_id", "signal_type", "headline", "summary",
        "local_report_decision", "selection_reason", "global_rank",
        "global_selection_reason", "global_report_summary", "editorial_note", "shortlist_lanes",
        "shortlist_lane_rank", "shortlist_rank_fields", "trend_claim", "today_conversation_count",
        "baseline_daily_counts", "source_candidate_ids", "source_topic_titles",
        "platform_counts", "appversion_counts", "feature_candidate_counts",
        "confidence", "needs_human_review", "artifact_reason", "feedback_link",
        "human_report_decision", "human_signal_type", "human_grouping_result", "human_notes",
    ]
    sheet.append(headers)
    for layer in ("main", "observe", "manual_review"):
        for row in selected.get(layer) or []:
            values = [
                layer, row.get("insight_id"), row.get("signal_type"), row.get("headline"), row.get("summary"),
                row.get("local_report_decision"), row.get("selection_reason"), row.get("global_rank"),
                row.get("global_selection_reason"), row.get("global_report_summary"), row.get("editorial_note"),
                _json_text(row.get("shortlist_lanes") or []), row.get("shortlist_lane_rank"),
                _json_text(row.get("shortlist_rank_fields") or {}),
                row.get("trend_claim"), row.get("today_conversation_count"),
                _json_text(row.get("baseline_daily_counts") or []),
                _json_text(row.get("source_candidate_ids") or []),
                _json_text(row.get("source_topic_titles") or []),
                _json_text(row.get("platform_counts") or {}),
                _json_text(row.get("appversion_counts") or {}),
                _json_text(row.get("feature_candidate_counts") or {}),
                row.get("confidence"), bool(row.get("needs_human_review")), row.get("artifact_reason"),
                "查看原反馈" if row.get("source_links") else "", "", "", "", "",
            ]
            sheet.append(values)
            if row.get("source_links"):
                _write_link(sheet.cell(sheet.max_row, headers.index("feedback_link") + 1), row["source_links"][0])


def _write_candidate_sheet(sheet, candidates: list[dict[str, Any]]) -> None:
    headers = [
        "candidate_id", "daily_topic_id", "stable_topic_id", "title", "description",
        "today_conversation_count", "baseline_daily_counts", "lifecycle_verdict",
        "recall_reasons", "candidate_band", "allowed_trend_claims", "feature_policy",
        "needs_review", "feedback_link",
    ]
    sheet.append(headers)
    for row in candidates:
        sheet.append([
            row.get("candidate_id"), row.get("daily_topic_id"), row.get("stable_topic_id"),
            row.get("title"), row.get("description"), row.get("today_conversation_count"),
            _json_text(row.get("baseline_daily_counts") or []), row.get("lifecycle_verdict"),
            _json_text(row.get("recall_reasons") or []), row.get("candidate_band"),
            _json_text(row.get("allowed_trend_claims") or []), _json_text(row.get("feature_policy") or {}),
            bool(row.get("needs_review")), "查看原反馈" if row.get("source_links") else "",
        ])
        if row.get("source_links"):
            _write_link(sheet.cell(sheet.max_row, headers.index("feedback_link") + 1), row["source_links"][0])


def _write_relation_sheet(sheet, relations: list[dict[str, Any]]) -> None:
    headers = ["left_candidate_id", "right_candidate_id", "relation_type", "cosine_similarity"]
    sheet.append(headers)
    for row in relations:
        sheet.append([row.get(header) for header in headers])


def _write_media_sheet(sheet, media_appendix: list[dict[str, Any]]) -> None:
    headers = ["summary", "has_media_evidence", "feedback_link", "human_notes"]
    sheet.append(headers)
    for row in media_appendix:
        links = _media_links(row)
        sheet.append([
            row.get("summary") or row.get("title") or "文本信息不足，需查看媒体",
            True,
            "查看原反馈" if links else "",
            "",
        ])
        if links:
            _write_link(sheet.cell(sheet.max_row, headers.index("feedback_link") + 1), links[0])


def _write_link(cell, url: str, label: str = "查看原反馈") -> None:
    cell.value = label
    cell.hyperlink = url
    cell.style = "Hyperlink"


def _format_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F6B5F")
        cell.alignment = Alignment(vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column_index in range(1, sheet.max_column + 1):
        values = [str(sheet.cell(row, column_index).value or "") for row in range(1, min(sheet.max_row, 30) + 1)]
        width = min(42, max(12, max((len(value) for value in values), default=0) + 2))
        sheet.column_dimensions[get_column_letter(column_index)].width = width


def _media_links(row: dict[str, Any]) -> list[str]:
    direct = str(row.get("evidence_link") or "").strip()
    values = [direct] if direct else []
    values.extend(row.get("evidence_links") or [])
    return _dedupe(values)


def _dedupe(values) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        text = str(value or "").strip()
        if text and text not in seen:
            seen.add(text)
            result.append(text)
    return result


def _json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True)
