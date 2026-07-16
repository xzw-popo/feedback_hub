from __future__ import annotations

import pytest
from openpyxl import load_workbook

from feedback_hub.topic_discovery import insight_artifacts as artifacts


def _candidate(candidate_id: str, *, links: list[str] | None = None) -> dict:
    return {
        "candidate_id": candidate_id,
        "daily_topic_id": f"daily:{candidate_id}",
        "stable_topic_id": f"topic:{candidate_id}",
        "title": f"候选{candidate_id}",
        "description": f"候选{candidate_id}描述",
        "today_conversation_ids": [f"conversation-{candidate_id}"],
        "today_conversation_count": 1,
        "baseline_dates": ["2026-07-07", "2026-07-08"],
        "baseline_daily_counts": [0, 1],
        "baseline_conversation_ids_by_date": {
            "2026-07-07": [],
            "2026-07-08": [f"history-{candidate_id}"],
        },
        "source_links": links if links is not None else [f"https://feedback/{candidate_id}"],
        "recall_reasons": ["high_value_single"],
        "candidate_band": "exploratory",
        "feedback_type_counts": {"bug_problem": 1},
        "feature_candidate_counts": {"voice_input": 1},
        "platform_counts": {"Android": 1},
        "appversion_counts": {"3.5.0": 1},
        "representative_issue_units": [{
            "issue_unit_id": f"issue-{candidate_id}",
            "summary": f"候选{candidate_id}",
            "evidence_spans": ["明确证据"],
        }],
        "known_context": True,
    }


def _insight(
    insight_id: str,
    decision: str = "nominate",
    *,
    confidence: float = 0.9,
    candidate_id: str | None = None,
) -> dict:
    candidate = candidate_id or insight_id
    return {
        "insight_id": insight_id,
        "report_decision": decision,
        "signal_type": "new_bug",
        "headline": f"洞察{insight_id}",
        "summary": f"洞察{insight_id}摘要",
        "selection_reason": "局部证据明确",
        "trend_claim": "none",
        "source_candidate_ids": [candidate],
        "representative_issue_unit_ids": [f"issue-{candidate}"],
        "confidence": confidence,
        "needs_human_review": decision == "manual_review",
    }


def _selection(*insight_ids: str) -> dict:
    return {
        "selected_insights": [{
            "insight_id": insight_id,
            "rank": index,
            "selection_reason": f"全局比较后选择{insight_id}",
            "editorial_note": "保持谨慎表达",
        } for index, insight_id in enumerate(insight_ids, 1)],
        "selection_summary": "完成全局比较",
    }


def _shortlist(*insight_ids: str) -> list[dict]:
    return [{
        "insight_id": insight_id,
        "shortlist_lanes": ["new_bug"],
        "shortlist_lane_rank": index,
        "shortlist_rank_fields": {"today_conversation_count": 1},
    } for index, insight_id in enumerate(insight_ids, 1)]


def test_partition_uses_only_ranked_global_selection_for_main() -> None:
    insights = [
        _insight("i1", candidate_id="c1"),
        _insight("i2", "observe", candidate_id="c2"),
        _insight("i3", candidate_id="c3"),
    ]
    candidates = {f"c{index}": _candidate(f"c{index}") for index in range(1, 4)}

    selected = artifacts.partition_selected_insights(
        insights,
        candidates,
        _selection("i2", "i1"),
        shortlist=_shortlist("i1", "i2"),
    )

    assert [row["insight_id"] for row in selected["main"]] == ["i2", "i1"]
    assert [row["insight_id"] for row in selected["observe"]] == ["i3"]
    assert selected["main"][0]["local_report_decision"] == "observe"
    assert selected["main"][0]["global_selection_reason"] == "全局比较后选择i2"
    assert selected["main"][0]["shortlist_lanes"] == ["new_bug"]


def test_partition_does_not_fill_unused_slots() -> None:
    selected = artifacts.partition_selected_insights(
        [_insight("i1", candidate_id="c1")],
        {"c1": _candidate("c1")},
        _selection(),
        shortlist=_shortlist("i1"),
    )

    assert selected["main"] == []
    assert selected["observe"][0]["artifact_reason"] == "not_selected_globally"


def test_partition_rejects_unknown_or_linkless_selected_insight() -> None:
    with pytest.raises(ValueError, match="unknown insight"):
        artifacts.partition_selected_insights(
            [_insight("i1", candidate_id="c1")],
            {"c1": _candidate("c1")},
            _selection("unknown"),
            shortlist=_shortlist("i1"),
        )

    with pytest.raises(ValueError, match="source link"):
        artifacts.partition_selected_insights(
            [_insight("i1", candidate_id="c1")],
            {"c1": _candidate("c1", links=[])},
            _selection("i1"),
            shortlist=_shortlist("i1"),
        )


def test_grouped_insight_deduplicates_conversations() -> None:
    left = _candidate("c1")
    right = _candidate("c2")
    right["today_conversation_ids"] = left["today_conversation_ids"] + ["conversation-c2"]
    insight = _insight("i1", candidate_id="c1")
    insight["source_candidate_ids"] = ["c1", "c2"]

    selected = artifacts.partition_selected_insights(
        [insight],
        {"c1": left, "c2": right},
        _selection("i1"),
        shortlist=_shortlist("i1"),
    )

    assert selected["main"][0]["today_conversation_count"] == 2


def test_report_contains_global_reason_sources_and_media_appendix() -> None:
    selected = artifacts.partition_selected_insights(
        [_insight("i1", candidate_id="c1")],
        {"c1": _candidate("c1", links=["https://feedback/main"])},
        _selection("i1"),
        shortlist=_shortlist("i1"),
    )
    text = artifacts.render_daily_report(
        "2026-07-14",
        selected,
        [{"summary": "图片展示异常", "evidence_link": "https://feedback/media"}],
        {"baseline_start": "2026-07-07", "baseline_end": "2026-07-13", "route_counts": {"glm": 2}},
    )

    assert "全局比较后选择i1" in text
    assert "[查看原反馈](https://feedback/main)" in text
    assert "媒体附录" in text
    assert "https://feedback/media" in text


def test_review_workbook_has_selection_audit_links_and_human_columns(tmp_path) -> None:
    candidate = _candidate("c1")
    selected = artifacts.partition_selected_insights(
        [_insight("i1", candidate_id="c1")],
        {"c1": candidate},
        _selection("i1"),
        shortlist=_shortlist("i1"),
    )
    path = tmp_path / "review.xlsx"

    artifacts.write_review_workbook(
        path,
        selected=selected,
        candidates=[candidate],
        relations=[{
            "left_candidate_id": "c1",
            "right_candidate_id": "c2",
            "relation_type": "semantic_similarity",
            "cosine_similarity": 0.82,
        }],
        media_appendix=[{"summary": "图片反馈", "evidence_link": "https://feedback/media"}],
    )

    book = load_workbook(path, data_only=False)
    assert book.sheetnames == ["正文与观察", "全部候选", "原子主题关系", "媒体附录"]
    headers = [cell.value for cell in book["正文与观察"][1]]
    for name in (
        "local_report_decision",
        "global_rank",
        "global_selection_reason",
        "editorial_note",
        "shortlist_lanes",
        "human_report_decision",
        "human_signal_type",
        "human_grouping_result",
        "human_notes",
    ):
        assert name in headers
    assert any(
        cell.hyperlink and cell.hyperlink.target.startswith("https://feedback/")
        for sheet in book.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )
