from __future__ import annotations

import hashlib
import json
import sqlite3
from datetime import datetime, timedelta, timezone

import openpyxl
from fastapi import FastAPI
from fastapi.testclient import TestClient

from feedback_hub.topic_mining.api import make_router
from feedback_hub.topic_mining.config import TopicMiningConfig
from feedback_hub.topic_mining.contracts import validate_topic_spec
from feedback_hub.topic_mining.export import HEADERS, export_topic_run
from feedback_hub.topic_mining.run_store import TopicRunStore
from feedback_hub.topic_mining.service import (
    get_candidate_page,
    run_topic_job,
    submit_caller_classifications,
    verify_topic_run,
)
from feedback_hub.topic_mining.vector_client import (
    VectorCapabilities,
    VectorHit,
    VectorSearchResult,
)


def _topic_spec(start: datetime, end: datetime):
    return validate_topic_spec({
        "schema_version": 1,
        "topic_name": "全屏时输入法工具栏不隐藏",
        "objective": "找出 Win 端游戏或视频全屏时输入法工具栏仍显示的反馈",
        "scope": {
            "start_time": start.isoformat(),
            "end_time": end.isoformat(),
            "platforms": ["Windows"],
            "products": ["微信输入法"],
        },
        "unit": "feedback",
        "inclusion_criteria": ["游戏或视频全屏时输入法工具栏仍显示"],
        "exclusion_criteria": ["Windows 系统任务栏仍显示", "游戏黑屏"],
        "positive_examples": ["游戏全屏时输入法工具栏一直显示"],
        "negative_examples": ["进入游戏后黑屏"],
        "lexical_hints": {"objects": ["工具栏"], "contexts": ["游戏", "全屏"]},
        "classification_labels": [
            {"id": "matched", "meaning": "明确符合专题定义"},
            {"id": "not_matched", "meaning": "不符合或证据不足"},
        ],
        "output": {
            "preferred_format": "xlsx",
            "required_fields": [
                "feedback_text", "feedback_time", "source_url",
            ],
        },
    })


def _create_source(path, start: datetime, end: datetime) -> None:
    start_ms = int(start.timestamp() * 1000)
    end_ms = int(end.timestamp() * 1000)
    with sqlite3.connect(path) as connection:
        connection.execute(
            """CREATE TABLE feedback (
                feedback_id TEXT PRIMARY KEY, conversation_id TEXT,
                msg_seq INTEGER, ts_ms INTEGER, platform TEXT,
                appversion TEXT, channel TEXT, device_name TEXT,
                user_vid TEXT, service_vid INTEGER,
                external_chat_url TEXT, text TEXT
            )"""
        )
        connection.executemany(
            "INSERT INTO feedback VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            [
                (
                    "win-game", "c-game", 1, start_ms, "Win", "1", "pc",
                    "PC", "u1", 1, "https://example.test/game",
                    "游戏全屏时输入法工具栏一直显示",
                ),
                (
                    "win-taskbar", "c-taskbar", 1, start_ms + 1, "Win",
                    "1", "pc", "PC", "u2", 2,
                    "https://example.test/taskbar",
                    "全屏时 Windows 系统任务栏仍显示",
                ),
                (
                    "win-black-screen", "c-black", 1, start_ms + 2,
                    "Win", "1", "pc", "PC", "u3", 3,
                    "https://example.test/black", "游戏进入全屏后黑屏",
                ),
                (
                    "mac-semantic", "c-mac", 1, start_ms + 3, "Mac",
                    "1", "pc", "Mac", "u4", 4,
                    "https://example.test/mac",
                    "视频全屏时输入法工具栏一直显示",
                ),
                (
                    "win-paraphrase", "c-vector", 1, start_ms + 4, "Win",
                    "1", "pc", "PC", "u5", 5,
                    "https://example.test/vector",
                    "沉浸场景悬浮控件遮住画面且始终没有收起",
                ),
                (
                    "coverage-boundary", "c-boundary", 1, end_ms, "Win",
                    "1", "pc", "PC", "u6", 6,
                    "https://example.test/boundary", "范围结束边界",
                ),
            ],
        )
        connection.execute(
            """CREATE TABLE feedback_source_coverage (
                channel TEXT NOT NULL, start_ts_ms INTEGER NOT NULL,
                end_ts_ms INTEGER NOT NULL, completed_at_ms INTEGER NOT NULL
            )"""
        )
        connection.execute(
            "INSERT INTO feedback_source_coverage VALUES (?, ?, ?, ?)",
            ("pc", start_ms, end_ms, end_ms),
        )


def test_caller_ai_end_to_end_partially_repairs_and_exports(tmp_path):
    start = datetime(2024, 1, 1, tzinfo=timezone.utc)
    end = start + timedelta(days=1)
    watermark = int(end.timestamp() * 1000)
    source = tmp_path / "source.db"
    _create_source(source, start, end)
    source_sha_before = hashlib.sha256(source.read_bytes()).hexdigest()
    spec = _topic_spec(start, end)
    assert spec.scope.platforms == ("Win",)
    config = TopicMiningConfig(
        source_db_path=source,
        data_dir=tmp_path / "topic-mining-data",
        vector_api_url="https://vector.example.test",
        vector_max_lag_seconds=1_000_000,
    )
    store = TopicRunStore(config.data_dir / "runs.db", config.data_dir / "runs")
    run = store.create_or_get(
        spec,
        watermark,
        classification_protocol_version=2,
        classification_owner="caller_ai",
    )

    class FakeVector:
        def capabilities(self):
            return VectorCapabilities(
                "feedback-items-v1", 1, ("feedback",), watermark,
            )

        def search(self, _queries, filters, _limit):
            assert filters["platforms"] == ["Win"]
            return VectorSearchResult(
                "feedback-items-v1",
                watermark,
                (
                    VectorHit("win-paraphrase", "objective:0", 0.99, 1),
                    VectorHit("mac-semantic", "objective:0", 0.98, 2),
                ),
            )

    outcome = run_topic_job(
        run["run_id"],
        store=store,
        config=config,
        vector_client=FakeVector(),
        model_call_fn=lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("backend model called"),
        ),
    )
    assert outcome["status"] == "classification_ready"
    artifact_dir = config.data_dir / "runs" / run["run_id"]
    assert not (artifact_dir / "classified.jsonl").exists()

    page = get_candidate_page(run["run_id"], 0, 20, store=store)
    assert {row["item_id"] for row in page["items"]} == {
        "win-game", "win-taskbar", "win-black-screen", "win-paraphrase",
    }
    paraphrase = next(
        row for row in page["items"] if row["item_id"] == "win-paraphrase"
    )
    assert paraphrase["channels"] == ["vector"]
    decisions = [
        {
            "item_id": "win-game", "label": "matched",
            "reason": "游戏全屏时工具栏仍显示",
            "evidence": ["输入法工具栏一直显示"],
        },
        {
            "item_id": "win-taskbar", "label": "not_matched",
            "reason": "明确排除系统任务栏", "evidence": [],
        },
        {
            "item_id": "win-black-screen", "label": "not_matched",
            "reason": "仅描述黑屏", "evidence": [],
        },
        {
            "item_id": "win-paraphrase", "label": "matched",
            "reason": "语义符合但先提交错误证据",
            "evidence": ["原文中不存在"],
        },
    ]
    partial = submit_caller_classifications(
        run["run_id"], decisions, store=store,
    )
    assert partial["accepted_count"] == 3
    assert partial["pending_count"] == 1
    assert partial["rejected"] == [{
        "item_id": "win-paraphrase",
        "code": "evidence_not_grounded",
    }]
    assert not (artifact_dir / "caller_classifications.jsonl").exists()

    repaired = submit_caller_classifications(
        run["run_id"],
        [{
            "item_id": "win-paraphrase", "label": "matched",
            "reason": "悬浮控件遮挡且未收起",
            "evidence": ["悬浮控件遮住画面"],
        }],
        store=store,
    )
    assert repaired["accepted_count"] == 4
    assert repaired["pending_count"] == 0
    assert store.get(run["run_id"])["status"] == "verification_ready"
    assert verify_topic_run(run["run_id"], store=store) == {
        "run_id": run["run_id"],
        "status": "verified",
        "matched_count": 2,
    }

    app = FastAPI()
    app.include_router(make_router(config=config, store=store))
    with TestClient(app) as client:
        public = client.get(
            f"/api/topic-mining/runs/{run['run_id']}",
        ).json()
    assert public["classification_owner"] == "caller_ai"
    assert public["accepted_decision_count"] == 4
    assert public["pending_decision_count"] == 0

    workbook_path = export_topic_run(run["run_id"], "xlsx", store=store)
    workbook = openpyxl.load_workbook(
        workbook_path, read_only=True, data_only=True,
    )
    sheet = workbook["反馈清单"]
    headers = [
        cell.value for cell in next(
            sheet.iter_rows(min_row=1, max_row=1),
        )
    ]
    feedback_id_column = headers.index("Feedback ID")
    workbook_ids = {
        row[feedback_id_column]
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[feedback_id_column]
    }
    workbook.close()
    assert headers == HEADERS
    assert workbook_ids == {"win-game", "win-paraphrase"}
    final_rows = [
        json.loads(line)
        for line in (
            artifact_dir / "final_reviewed.jsonl"
        ).read_text(encoding="utf-8").splitlines()
    ]
    assert all(row["source"] == "caller_ai" for row in final_rows)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_sha_before
