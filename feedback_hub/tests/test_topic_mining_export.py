from __future__ import annotations

import json
import sqlite3
from pathlib import Path

import openpyxl
import pytest
import hashlib

from feedback_hub.topic_mining.contracts import validate_topic_spec
from feedback_hub.topic_mining.export import export_topic_run
from feedback_hub.jsonl_io import load_jsonl_objects
from feedback_hub.topic_mining.run_store import TopicRunStore
from feedback_hub.topic_mining.service import RunVerificationError

CUTOFF_MS = 1_700_000_000_000


def _spec(*, mode: str = "standard"):
    return validate_topic_spec({"schema_version": 1, "topic_name": "专题", "objective": "找工具栏", "scope": {"start_time": "2023-11-14T00:00:00+00:00", "end_time": "2023-11-16T00:00:00+00:00", "platforms": ["Win"], "products": ["微信输入法"]}, "unit": "feedback", "mode": mode, "inclusion_criteria": ["工具栏"], "exclusion_criteria": ["系统任务栏"], "positive_examples": ["全屏工具栏"], "negative_examples": ["黑屏"], "lexical_hints": {"objects": ["工具栏"]}, "classification_labels": [{"id": "matched", "meaning": "命中"}, {"id": "not_matched", "meaning": "不命中"}], "output": {"preferred_format": "xlsx", "required_fields": ["feedback_text"]}})


def _row(run_id: str, item_id: str = "f-1") -> dict:
    return {
        "item_id": item_id, "label": "matched", "confidence": 0.9,
        "evidence": ["工具栏一直显示"], "reason": "全屏仍显示", "source": "classifier",
        "run_id": run_id, "data_cutoff_ms": CUTOFF_MS,
        "source_item": {
            "item_id": item_id, "feedback_id": item_id, "conversation_id": "c-1",
            "ts_ms": CUTOFF_MS, "platform": "Win", "appversion": "1.2",
            "channel": "pc", "device_name": "PC", "text": "游戏全屏工具栏一直显示",
            "source_url": "https://example.test/chat/1",
        },
    }


def _rows(run_id: str, count: int) -> list[dict]:
    rows = []
    for index in range(count):
        row = _row(run_id, f"f-{index:03d}")
        item = row["source_item"]
        item.update({
            "ts_ms": CUTOFF_MS - index,
            "channel": "pc" if index % 2 else "mobile",
            "text": f"游戏全屏工具栏一直显示 {index}",
            "source_url": f"https://example.test/chat/{index}",
        })
        row["evidence"] = [f"工具栏一直显示 {index}"]
        rows.append(row)
    return rows


def _read_jsonl(path: Path) -> list[dict]:
    return [
        json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]


def _persist_verified_manifest(
    store: TopicRunStore,
    run: dict,
    *files: Path,
) -> dict:
    artifact_dir = Path(run["artifact_dir"])
    snapshot = artifact_dir / "source_snapshot.sqlite"
    with sqlite3.connect(snapshot) as connection:
        connection.execute("CREATE TABLE feedback (ts_ms INTEGER NOT NULL)")
        connection.execute(
            "INSERT INTO feedback VALUES (?)", (run["source_watermark_ms"],),
        )
        connection.execute(
            """CREATE TABLE feedback_source_coverage (
                completed_at_ms INTEGER NOT NULL
            )"""
        )
        connection.execute(
            "INSERT INTO feedback_source_coverage VALUES (?)",
            (run["source_watermark_ms"],),
        )
    snapshot_digest = hashlib.sha256(snapshot.read_bytes()).hexdigest()
    manifest = {
        "source_watermark_ms": run["source_watermark_ms"],
        "source_snapshot": {
            "max_ts_ms": run["source_watermark_ms"],
            "coverage_watermark_ms": run["source_watermark_ms"],
            "sha256": snapshot_digest,
        },
        "source_sha256": snapshot_digest,
        "artifacts": {
            snapshot.name: snapshot_digest,
            **{
                path.name: hashlib.sha256(path.read_bytes()).hexdigest()
                for path in files
            },
        },
    }
    store.update_manifest(
        run["run_id"], manifest, stage="verified", status="verified",
    )
    (artifact_dir / "manifest.json").write_text(
        json.dumps(manifest), encoding="utf-8",
    )
    return manifest


@pytest.mark.parametrize("export_format", ["jsonl", "xlsx"])
def test_export_preserves_unicode_separator_in_recall_pool(tmp_path, export_format):
    store = TopicRunStore(tmp_path / "runs.db", tmp_path / "runs")
    run = store.create_or_get(_spec(), CUTOFF_MS)
    artifact_dir = Path(run["artifact_dir"])
    final = artifact_dir / "final_reviewed.jsonl"
    final.write_text(
        json.dumps(_row(run["run_id"]), ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    recall_rows = []
    for item_id, text in (
        ("f-1", "游戏全屏工具栏一直显示"),
        ("f-2", "第一段\u2028第二段"),
    ):
        item = dict(_row(run["run_id"], item_id)["source_item"])
        item["text"] = text
        recall_rows.append({
            "item_id": item_id, "item": item, "channels": ["bm25"],
            "fused_score": 1.0, "fused_rank": len(recall_rows) + 1,
            "channel_ranks": {"bm25": len(recall_rows) + 1},
            "raw_scores": {"bm25": 1.0}, "query_ids": ["q1"],
            "negative_query_hits": [],
        })
    recalls = artifact_dir / "recall_candidates.jsonl"
    recalls.write_text(
        "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in recall_rows),
        encoding="utf-8",
    )
    _persist_verified_manifest(store, run, final, recalls)

    exported = export_topic_run(run["run_id"], export_format, store=store)

    assert exported.is_file()
    persisted = load_jsonl_objects(recalls.read_bytes())
    assert persisted[1]["item"]["text"] == "第一段\u2028第二段"


def test_export_has_unique_ids_links_and_evidence(tmp_path):
    store = TopicRunStore(tmp_path / "runs.db", tmp_path / "runs")
    run = store.create_or_get(_spec(), CUTOFF_MS)
    artifact_dir = tmp_path / "runs" / run["run_id"]
    (artifact_dir / "final_reviewed.jsonl").write_text(json.dumps(_row(run["run_id"])) + "\n", encoding="utf-8")
    _persist_verified_manifest(
        store, run, artifact_dir / "final_reviewed.jsonl",
    )
    path = export_topic_run(run["run_id"], "xlsx", store=store)
    wb = openpyxl.load_workbook(path, read_only=False, data_only=False)
    ws = wb["反馈清单"]
    headers = [cell.value for cell in ws[1]]
    assert headers == [
        "反馈时间", "反馈原文", "对应链接", "平台", "版本", "设备",
        "Feedback ID", "判定理由", "证据",
    ]
    assert not {
        "命中分类", "Conversation ID", "Run ID", "数据截止时间",
    } & set(headers)
    assert ws["C2"].hyperlink.target.startswith("https://")
    assert ws.cell(2, headers.index("判定理由") + 1).value == "全屏仍显示"
    assert ws.cell(2, headers.index("证据") + 1).value == "工具栏一直显示"
    assert (artifact_dir / "final_results.jsonl").exists()
    assert (artifact_dir / "quality_report.json").exists()
    final_rows = [json.loads(line) for line in (artifact_dir / "final_results.jsonl").read_text(encoding="utf-8").splitlines()]
    assert [row["data_cutoff_ms"] for row in final_rows] == [CUTOFF_MS]
    assert [
        {
            "feedback_text": row["feedback_text"],
            "feedback_time": row["feedback_time"],
            "source_url": row["source_url"],
        }
        for row in final_rows
    ] == [{
        "feedback_text": "游戏全屏工具栏一直显示",
        "feedback_time": "2023-11-14T22:13:20+00:00",
        "source_url": "https://example.test/chat/1",
    }]
    report = json.loads((artifact_dir / "quality_report.json").read_text(encoding="utf-8"))
    assert report["data_cutoff_ms"] == CUTOFF_MS
    assert report["required_fields"] == ["feedback_text"]


def test_standard_export_keeps_all_321_confirmed_matches(tmp_path):
    store = TopicRunStore(tmp_path / "runs.db", tmp_path / "runs")
    run = store.create_or_get(_spec(), CUTOFF_MS)
    artifact_dir = Path(run["artifact_dir"])
    final = artifact_dir / "final_reviewed.jsonl"
    final.write_text(
        "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in _rows(run["run_id"], 321)),
        encoding="utf-8",
    )
    verified = _persist_verified_manifest(store, run, final)
    verified.update({"retrieved_candidate_count": 570, "classified_count": 500})
    store.update_manifest(run["run_id"], verified, stage="verified", status="verified")
    (artifact_dir / "manifest.json").write_text(json.dumps(verified), encoding="utf-8")

    jsonl_path = export_topic_run(run["run_id"], "jsonl", store=store)
    xlsx_path = export_topic_run(run["run_id"], "xlsx", store=store)

    assert len(_read_jsonl(jsonl_path)) == 321
    workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    assert workbook["反馈清单"].max_row == 322
    metadata = dict(workbook["导出元数据"].iter_rows(values_only=True))
    workbook.close()
    assert metadata == {
        "mode": "standard", "result_scope": "representative",
        "matched_total": 321, "returned_feedback": 321,
        "possibly_more_matches": True,
    }
    report = json.loads((artifact_dir / "quality_report.json").read_text(encoding="utf-8"))
    assert {key: report[key] for key in ("mode", "result_scope", "matched_total", "returned_feedback", "possibly_more_matches")} == {
        "mode": "standard", "result_scope": "representative", "matched_total": 321,
        "returned_feedback": 321, "possibly_more_matches": True,
    }
    manifest = json.loads(store.get(run["run_id"])["manifest_json"])
    assert manifest["result_scope"] == "representative"
    assert len(_read_jsonl(final)) == 321


def test_standard_export_under_limit_reports_possible_unclassified_matches_in_xlsx_metadata(tmp_path):
    store = TopicRunStore(tmp_path / "runs.db", tmp_path / "runs")
    run = store.create_or_get(_spec(), CUTOFF_MS)
    artifact_dir = Path(run["artifact_dir"])
    final = artifact_dir / "final_reviewed.jsonl"
    final.write_text(
        "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in _rows(run["run_id"], 23)),
        encoding="utf-8",
    )
    manifest = _persist_verified_manifest(store, run, final)
    manifest.update({"retrieved_candidate_count": 130, "classified_count": 23})
    store.update_manifest(run["run_id"], manifest, stage="verified", status="verified")
    (artifact_dir / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")

    path = export_topic_run(run["run_id"], "xlsx", store=store)

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    metadata = dict(workbook["导出元数据"].iter_rows(values_only=True))
    workbook.close()
    assert metadata == {
        "mode": "standard", "result_scope": "representative", "matched_total": 23,
        "returned_feedback": 23, "possibly_more_matches": True,
    }


def test_exhaustive_export_keeps_all_reviewed_matches(tmp_path):
    store = TopicRunStore(tmp_path / "runs.db", tmp_path / "runs")
    run = store.create_or_get(_spec(mode="exhaustive"), CUTOFF_MS)
    artifact_dir = Path(run["artifact_dir"])
    final = artifact_dir / "final_reviewed.jsonl"
    final.write_text(
        "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in _rows(run["run_id"], 130)),
        encoding="utf-8",
    )
    _persist_verified_manifest(store, run, final)

    path = export_topic_run(run["run_id"], "jsonl", store=store)

    assert len(_read_jsonl(path)) == 130
    report = json.loads((artifact_dir / "quality_report.json").read_text(encoding="utf-8"))
    assert {key: report[key] for key in ("mode", "result_scope", "matched_total", "returned_feedback", "possibly_more_matches")} == {
        "mode": "exhaustive", "result_scope": "reviewed", "matched_total": 130,
        "returned_feedback": 130, "possibly_more_matches": False,
    }


def test_exhaustive_export_reports_possible_matches_beyond_5000_candidate_safety_limit(tmp_path):
    store = TopicRunStore(tmp_path / "runs.db", tmp_path / "runs")
    run = store.create_or_get(_spec(mode="exhaustive"), CUTOFF_MS)
    artifact_dir = Path(run["artifact_dir"])
    final = artifact_dir / "final_reviewed.jsonl"
    final.write_text(
        "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in _rows(run["run_id"], 130)),
        encoding="utf-8",
    )
    manifest = _persist_verified_manifest(store, run, final)
    manifest.update({"retrieved_candidate_count": 5_001, "classified_count": 5_000})
    store.update_manifest(run["run_id"], manifest, stage="verified", status="verified")
    (artifact_dir / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")

    path = export_topic_run(run["run_id"], "jsonl", store=store)

    assert len(_read_jsonl(path)) == 130
    report = json.loads((artifact_dir / "quality_report.json").read_text(encoding="utf-8"))
    assert report["mode"] == "exhaustive"
    assert report["result_scope"] == "representative"
    assert report["possibly_more_matches"] is True


def test_export_uses_terminal_manifest_cas(tmp_path, monkeypatch):
    store = TopicRunStore(tmp_path / "runs.db", tmp_path / "runs")
    run = store.create_or_get(_spec(), CUTOFF_MS)
    artifact_dir = Path(run["artifact_dir"])
    final = artifact_dir / "final_reviewed.jsonl"
    final.write_text(json.dumps(_row(run["run_id"])) + "\n", encoding="utf-8")
    _persist_verified_manifest(store, run, final)
    before = store.get(run["run_id"])
    publications = []
    publish = store.publish_terminal_artifacts

    def observe(*args, **kwargs):
        publications.append(kwargs)
        return publish(*args, **kwargs)

    monkeypatch.setattr(store, "publish_terminal_artifacts", observe)

    export_topic_run(run["run_id"], "jsonl", store=store)

    assert len(publications) == 1
    assert publications[0]["expected_manifest_json"] == before["manifest_json"]
    assert set(publications[0]["files"]) == {
        "final_results.jsonl", "quality_report.json",
    }
    current_manifest = json.loads(store.get(run["run_id"])["manifest_json"])
    assert set(current_manifest["artifacts"]) == {
        "source_snapshot.sqlite", "final_reviewed.jsonl",
        "final_results.jsonl", "quality_report.json",
    }


def test_export_rejects_a_run_whose_persisted_identity_was_tampered(tmp_path):
    store = TopicRunStore(tmp_path / "runs.db", tmp_path / "runs")
    run = store.create_or_get(_spec(), CUTOFF_MS)
    artifact_dir = Path(run["artifact_dir"])
    final = artifact_dir / "final_reviewed.jsonl"
    final.write_text(json.dumps(_row(run["run_id"])) + "\n", encoding="utf-8")
    _persist_verified_manifest(store, run, final)
    with sqlite3.connect(store.db_path) as connection:
        connection.execute(
            "UPDATE topic_run SET spec_hash = 'forged' WHERE run_id = ?",
            (run["run_id"],),
        )

    with pytest.raises(RunVerificationError, match="run_identity_mismatch"):
        export_topic_run(run["run_id"], "jsonl", store=store)


def test_export_requires_complete_frozen_snapshot_evidence(tmp_path):
    store = TopicRunStore(tmp_path / "runs.db", tmp_path / "runs")
    run = store.create_or_get(_spec(), CUTOFF_MS)
    artifact_dir = Path(run["artifact_dir"])
    final = artifact_dir / "final_reviewed.jsonl"
    final.write_text(json.dumps(_row(run["run_id"])) + "\n", encoding="utf-8")
    manifest = {
        "artifacts": {final.name: hashlib.sha256(final.read_bytes()).hexdigest()},
    }
    store.update_manifest(
        run["run_id"], manifest, stage="verified", status="verified",
    )
    (artifact_dir / "manifest.json").write_text(
        json.dumps(manifest), encoding="utf-8",
    )

    with pytest.raises(RunVerificationError, match="source_snapshot_required"):
        export_topic_run(run["run_id"], "jsonl", store=store)


@pytest.mark.parametrize("cutoff", [None, True, CUTOFF_MS - 1, str(CUTOFF_MS)])
def test_export_rejects_missing_invalid_or_mismatched_data_cutoff(tmp_path, cutoff):
    store = TopicRunStore(tmp_path / "runs.db", tmp_path / "runs")
    run = store.create_or_get(_spec(), CUTOFF_MS)
    artifact_dir = tmp_path / "runs" / run["run_id"]
    row = _row(run["run_id"])
    if cutoff is None:
        row.pop("data_cutoff_ms")
    else:
        row["data_cutoff_ms"] = cutoff
    final = artifact_dir / "final_reviewed.jsonl"
    final.write_text(json.dumps(row) + "\n", encoding="utf-8")
    _persist_verified_manifest(store, run, final)

    with pytest.raises(RunVerificationError, match="data_cutoff"):
        export_topic_run(run["run_id"], "jsonl", store=store)


def test_export_rejects_item_newer_than_advertised_cutoff(tmp_path):
    store = TopicRunStore(tmp_path / "runs.db", tmp_path / "runs")
    run = store.create_or_get(_spec(), CUTOFF_MS)
    artifact_dir = Path(run["artifact_dir"])
    row = _row(run["run_id"])
    row["source_item"]["ts_ms"] = CUTOFF_MS + 1
    final = artifact_dir / "final_reviewed.jsonl"
    final.write_text(json.dumps(row) + "\n", encoding="utf-8")
    _persist_verified_manifest(store, run, final)

    with pytest.raises(RunVerificationError, match="data_cutoff"):
        export_topic_run(run["run_id"], "jsonl", store=store)


@pytest.mark.parametrize("text", ["=HYPERLINK(\"bad\")", "+1+1", "-1+1", "@SUM(A1)"])
def test_export_escapes_formula_like_user_text(tmp_path, text):
    store = TopicRunStore(tmp_path / "runs.db", tmp_path / "runs")
    run = store.create_or_get(_spec(), CUTOFF_MS)
    artifact_dir = tmp_path / "runs" / run["run_id"]
    row = _row(run["run_id"])
    row["source_item"]["text"] = text
    row["evidence"] = [text]
    (artifact_dir / "final_reviewed.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")
    _persist_verified_manifest(
        store, run, artifact_dir / "final_reviewed.jsonl",
    )
    path = export_topic_run(run["run_id"], "xlsx", store=store)
    wb = openpyxl.load_workbook(path, data_only=False)
    assert wb["反馈清单"]["B2"].value.startswith("'")


def test_export_blocks_unverified_run(tmp_path):
    store = TopicRunStore(tmp_path / "runs.db", tmp_path / "runs")
    run = store.create_or_get(_spec(), 123)
    with pytest.raises(RunVerificationError, match="run_not_verified"):
        export_topic_run(run["run_id"], "jsonl", store=store)


def test_export_rejects_tampered_final_artifact(tmp_path):
    store = TopicRunStore(tmp_path / "runs.db", tmp_path / "runs")
    run = store.create_or_get(_spec(), 123)
    artifact_dir = tmp_path / "runs" / run["run_id"]
    final = artifact_dir / "final_reviewed.jsonl"
    final.write_text(json.dumps(_row(run["run_id"])) + "\n", encoding="utf-8")
    _persist_verified_manifest(store, run, final)
    final.write_text(json.dumps({**_row(run["run_id"]), "reason": "tampered"}) + "\n", encoding="utf-8")
    with pytest.raises(RunVerificationError, match="manifest_hash_reconciliation"):
        export_topic_run(run["run_id"], "xlsx", store=store)
