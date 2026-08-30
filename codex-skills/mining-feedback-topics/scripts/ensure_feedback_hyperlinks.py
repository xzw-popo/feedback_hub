#!/usr/bin/env python3
"""Repair and validate clickable links in feedback-topic Excel workbooks."""

from __future__ import annotations

import argparse
import json
import os
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet


ID_HEADERS = frozenset({"feedback id", "feedback_id", "反馈id"})
LINK_HEADERS = frozenset({"对应链接", "反馈链接", "source_url", "链接"})
FEEDBACK_HEADERS = frozenset({"反馈原文", "feedback_text"})
ALL_FEEDBACK_HEADERS = ID_HEADERS | LINK_HEADERS | FEEDBACK_HEADERS
MAX_HEADER_SCAN_ROWS = 20


@dataclass(frozen=True)
class ItemSheet:
    sheet: Worksheet
    header_row: int
    id_column: int
    link_column: int


def _normalized_header(value: object) -> str:
    if not isinstance(value, str):
        return ""
    return " ".join(value.strip().lower().split())


def _normalized_id(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _row_is_empty(sheet: Worksheet, row: int) -> bool:
    return all(
        cell.value is None
        or (isinstance(cell.value, str) and not cell.value.strip())
        for cell in sheet[row]
    )


def _find_item_sheets(workbook) -> list[ItemSheet]:
    found: list[ItemSheet] = []
    for sheet in workbook.worksheets:
        non_empty_scanned = 0
        for row in range(1, sheet.max_row + 1):
            values = [_normalized_header(cell.value) for cell in sheet[row]]
            if not any(values):
                continue
            non_empty_scanned += 1
            recognized = {
                index: value
                for index, value in enumerate(values, start=1)
                if value in ALL_FEEDBACK_HEADERS
            }
            if recognized:
                id_columns = [
                    index for index, value in recognized.items()
                    if value in ID_HEADERS
                ]
                link_columns = [
                    index for index, value in recognized.items()
                    if value in LINK_HEADERS
                ]
                if len(id_columns) != 1 or len(link_columns) != 1:
                    raise ValueError(
                        f"sheet {sheet.title!r} requires Feedback ID and link columns"
                    )
                found.append(ItemSheet(
                    sheet=sheet,
                    header_row=row,
                    id_column=id_columns[0],
                    link_column=link_columns[0],
                ))
                break
            if non_empty_scanned >= MAX_HEADER_SCAN_ROWS:
                break
    return found


def _expected_formula(target: str) -> str:
    escaped = target.replace('"', '""')
    return f'=HYPERLINK("{escaped}","打开反馈")'


def _validate_link_cell(cell: Cell, target: str) -> None:
    hyperlink_target = cell.hyperlink.target if cell.hyperlink is not None else None
    if (
        cell.data_type != "f"
        or cell.value != _expected_formula(target)
        or hyperlink_target != target
        or cell.style != "Hyperlink"
    ):
        raise ValueError(
            f"sheet {cell.parent.title!r} row {cell.row} has a plain-text "
            "or invalid feedback link"
        )


def _load_source_links(path: Path) -> dict[str, str]:
    try:
        workbook = load_workbook(path, data_only=False, read_only=False)
    except (OSError, ValueError) as error:
        raise ValueError(f"cannot read official workbook: {error}") from error
    sheets = _find_item_sheets(workbook)
    if not sheets:
        raise ValueError("official workbook has no item-level feedback sheet")
    links: dict[str, str] = {}
    for item_sheet in sheets:
        sheet = item_sheet.sheet
        for row in range(item_sheet.header_row + 1, sheet.max_row + 1):
            if _row_is_empty(sheet, row):
                continue
            feedback_id = _normalized_id(
                sheet.cell(row, item_sheet.id_column).value
            )
            if not feedback_id:
                raise ValueError("official feedback IDs must be non-empty")
            if feedback_id in links:
                raise ValueError("official feedback IDs must be unique")
            cell = sheet.cell(row, item_sheet.link_column)
            target = cell.hyperlink.target if cell.hyperlink is not None else None
            if not isinstance(target, str) or not target.startswith("https://"):
                raise ValueError(
                    f"feedback ID {feedback_id!r} has no valid official link"
                )
            _validate_link_cell(cell, target)
            links[feedback_id] = target
    return links


def _process_workbook(
    workbook,
    source_links: dict[str, str],
    *,
    repair: bool,
) -> dict[str, int]:
    sheets = _find_item_sheets(workbook)
    if not sheets:
        raise ValueError("workbook has no item-level feedback sheet")
    seen: set[str] = set()
    repaired = 0
    for item_sheet in sheets:
        sheet = item_sheet.sheet
        for row in range(item_sheet.header_row + 1, sheet.max_row + 1):
            if _row_is_empty(sheet, row):
                continue
            feedback_id = _normalized_id(
                sheet.cell(row, item_sheet.id_column).value
            )
            if not feedback_id:
                raise ValueError("feedback IDs must be non-empty")
            if feedback_id in seen:
                raise ValueError("feedback IDs must be unique")
            target = source_links.get(feedback_id)
            if target is None:
                raise ValueError(f"feedback ID {feedback_id!r} is unknown")
            seen.add(feedback_id)
            cell = sheet.cell(row, item_sheet.link_column)
            if repair:
                cell.value = _expected_formula(target)
                cell.hyperlink = target
                cell.style = "Hyperlink"
                repaired += 1
            _validate_link_cell(cell, target)
    return {
        "feedback_rows": len(seen),
        "item_sheets": len(sheets),
        "repaired_links": repaired,
    }


def _load_input(path: Path):
    try:
        return load_workbook(path, data_only=False, read_only=False)
    except (OSError, ValueError) as error:
        raise ValueError(f"cannot read input workbook: {error}") from error


def _save_atomic(workbook, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    handle = tempfile.NamedTemporaryFile(
        prefix=f".{output_path.stem}.",
        suffix=".tmp.xlsx",
        dir=output_path.parent,
        delete=False,
    )
    temporary = Path(handle.name)
    handle.close()
    try:
        workbook.save(temporary)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    return temporary


def repair_and_validate(
    official_path: Path,
    input_path: Path,
    output_path: Path | None,
) -> dict[str, int]:
    official_path = official_path.resolve()
    input_path = input_path.resolve()
    if official_path.suffix.lower() != ".xlsx":
        raise ValueError("official workbook must be an .xlsx file")
    if input_path.suffix.lower() != ".xlsx":
        raise ValueError("input workbook must be an .xlsx file")
    if output_path is not None:
        output_path = output_path.resolve()
        if output_path.suffix.lower() != ".xlsx":
            raise ValueError("output workbook must be an .xlsx file")
        if output_path in {input_path, official_path}:
            raise ValueError("output must differ from input and official workbooks")

    source_links = _load_source_links(official_path)
    workbook = _load_input(input_path)
    summary = _process_workbook(
        workbook,
        source_links,
        repair=output_path is not None,
    )
    if output_path is None:
        return summary

    temporary = _save_atomic(workbook, output_path)
    try:
        saved = _load_input(temporary)
        verified = _process_workbook(saved, source_links, repair=False)
        if verified["feedback_rows"] != summary["feedback_rows"]:
            raise ValueError("saved workbook feedback coverage changed")
        os.replace(temporary, output_path)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    return summary


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--official", required=True)
    parser.add_argument("--input", required=True)
    action = parser.add_mutually_exclusive_group(required=True)
    action.add_argument("--check", action="store_true")
    action.add_argument("--output")
    return parser


def main(arguments: list[str] | None = None) -> None:
    parser = _parser()
    try:
        parsed = parser.parse_args(arguments)
        summary = repair_and_validate(
            Path(parsed.official),
            Path(parsed.input),
            None if parsed.check else Path(parsed.output),
        )
    except (OSError, ValueError) as error:
        print(str(error), file=sys.stderr)
        raise SystemExit(2) from error
    print(json.dumps(summary, ensure_ascii=False, sort_keys=True))


if __name__ == "__main__":
    main()
